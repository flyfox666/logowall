import io
import json
import zipfile
import threading
from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook

from conftest import png_bytes


# ---- Clients ----------------------------------------------------------------
def test_client_crud_roundtrip(ctx):
    a = ctx.admin()
    r = ctx.client.post('/api/clients', headers=a, json={
        'company': 'Acme', 'office_code': 'sha', 'departments': '审计；税务', 'owners': '张三;李四',
        'cooperation_date': '2024/3/5', 'website': 'acme.com'})
    assert r.status_code == 200, r.text
    rec = r.json()
    assert rec['office_city'] == '上海' and rec['region'] == '华东'
    assert rec['departments'] == ['审计', '税务'] and rec['owners'] == ['张三', '李四']
    assert rec['cooperation_date'] == '2024-03-05'
    r = ctx.client.put(f"/api/clients/{rec['id']}", headers=a, json={'company': 'Acme 2'})
    assert r.json()['brand'] == 'Acme 2'
    assert ctx.client.get('/api/filters', headers=a).json()['owners'] == ['张三', '李四']
    assert ctx.client.delete(f"/api/clients/{rec['id']}", headers=a).status_code == 200
    assert ctx.client.get('/api/clients', headers=a).json()['total'] == 0


@pytest.mark.parametrize('field,value', [
    ('logo_url', 'javascript:alert(1)'),
    ('logo_url', 'logos/../users.json'),
    ('logo_url', 'data:image/svg+xml,<svg onload=alert(1)>'),
    ('website', 'javascript:alert(1)'),
])
def test_dangerous_urls_rejected(ctx, field, value):
    r = ctx.client.post('/api/clients', headers=ctx.admin(), json={'company': 'X', field: value})
    assert r.status_code == 400


def test_concurrent_writes_do_not_lose_records(ctx):
    a = ctx.admin()
    errors = []

    def worker(i):
        r = ctx.client.post('/api/clients', headers=a, json={'company': f'C{i}'})
        if r.status_code != 200:
            errors.append(r.text)

    threads = [threading.Thread(target=worker, args=(i,)) for i in range(25)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert not errors
    recs = ctx.client.get('/api/clients', headers=a).json()['records']
    assert len(recs) == 25
    assert len({r['id'] for r in recs}) == 25


def test_data_json_etag(ctx):
    a = ctx.admin()
    r = ctx.client.get('/data.json', headers=a)
    etag = r.headers['etag']
    r2 = ctx.client.get('/data.json', headers={**a, 'If-None-Match': etag})
    assert r2.status_code == 304
    ctx.client.post('/api/clients', headers=a, json={'company': 'Changed'})
    r3 = ctx.client.get('/data.json', headers={**a, 'If-None-Match': etag})
    assert r3.status_code == 200 and r3.json()['total_count'] == 1


def test_settings_validation(ctx):
    a = ctx.admin()
    assert ctx.client.put('/api/settings', headers=a, json={'title': 'Hi', 'custom_primary': '#aabbcc'}).status_code == 200
    assert ctx.client.put('/api/settings', headers=a, json={'custom_primary': 'red;}'}).status_code == 400
    assert ctx.client.put('/api/settings', headers=a, json={'title': 'x' * 1000}).status_code == 400


def test_health_reports_writable(ctx):
    assert ctx.client.get('/api/health').json() == {'status': 'ok', 'data_writable': True}


# ---- Logos ------------------------------------------------------------------
def test_upload_rejects_non_images_even_with_image_name(ctx):
    r = ctx.client.post('/api/logo/upload', headers=ctx.admin(),
                        files={'file': ('logo.png', b'<html><script>alert(1)</script>', 'image/png')})
    assert r.status_code == 400


def test_upload_png_and_serve_with_sandbox_csp(ctx):
    a = ctx.admin()
    r = ctx.client.post('/api/logo/upload', headers=a, files={'file': ('x.bin', png_bytes(), 'application/octet-stream')})
    assert r.status_code == 200, r.text
    url = r.json()['url']
    assert url.endswith('.png')  # type comes from content, not name
    served = ctx.client.get('/' + url, headers=a)
    assert served.status_code == 200
    assert 'sandbox' in served.headers['content-security-policy']


def test_svg_upload_is_sandboxed(ctx):
    a = ctx.admin()
    svg = b'<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script></svg>'
    r = ctx.client.post('/api/logo/upload', headers=a, files={'file': ('x.svg', svg, 'image/svg+xml')})
    url = r.json()['url']
    served = ctx.client.get('/' + url, headers=a)
    assert served.headers['content-type'].startswith('image/svg+xml')
    assert "default-src 'none'" in served.headers['content-security-policy']


def test_large_logo_is_downscaled(ctx):
    from PIL import Image
    a = ctx.admin()
    big = png_bytes(size=(3000, 1500))
    r = ctx.client.post('/api/logo/upload', headers=a, files={'file': ('big.png', big, 'image/png')})
    path = ctx.data_dir / r.json()['url']
    with Image.open(path) as im:
        assert max(im.size) == 800 and im.size == (800, 400)
        assert im.mode == 'RGBA'


def test_upload_size_limit(make_ctx):
    c = make_ctx(seed_records=[], MAX_UPLOAD_MB='1')
    r = c.client.post('/api/logo/upload', headers=c.admin(),
                      files={'file': ('x.png', b'\x89PNG\r\n\x1a\n' + b'0' * (1024 * 1024 + 10), 'image/png')})
    assert r.status_code == 413


def test_batch_upload_mixed(ctx):
    r = ctx.client.post('/api/logo/batch-upload', headers=ctx.admin(), files=[
        ('files', ('a.png', png_bytes(), 'image/png')),
        ('files', ('b.png', b'not an image', 'image/png')),
    ])
    body = r.json()
    assert body['uploaded'] == 1
    assert [x['ok'] for x in body['results']] == [True, False]


def test_replace_logo_changing_type_updates_clients(ctx):
    a = ctx.admin()
    url = ctx.client.post('/api/logo/upload', headers=a,
                          files={'file': ('a.png', png_bytes(), 'image/png')}).json()['url']
    cid = ctx.client.post('/api/clients', headers=a, json={'company': 'X', 'logo_url': url}).json()['id']
    svg = b'<svg xmlns="http://www.w3.org/2000/svg"></svg>'
    r = ctx.client.post('/api/logo/replace', headers=a, data={'filename': url.split('/')[1]},
                        files={'file': ('n.svg', svg, 'image/svg+xml')})
    assert r.status_code == 200, r.text
    new_url = r.json()['url']
    assert new_url.endswith('.svg')
    assert ctx.client.get(f'/api/clients/{cid}', headers=a).json()['logo_url'] == new_url
    assert not (ctx.data_dir / url).exists()


def test_logo_delete_path_traversal_and_in_use(ctx):
    a = ctx.admin()
    assert ctx.client.delete('/api/logos/..%2Fusers.json', headers=a).status_code in (400, 404)
    assert (ctx.data_dir / 'users.json').exists()
    url = ctx.client.post('/api/logo/upload', headers=a,
                          files={'file': ('a.png', png_bytes(), 'image/png')}).json()['url']
    ctx.client.post('/api/clients', headers=a, json={'company': 'X', 'logo_url': url})
    assert ctx.client.delete('/api/' + url.replace('logos/', 'logos/'), headers=a).status_code == 400


# ---- Backup -------------------------------------------------------------------
def _zip(entries):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zf:
        for name, data in entries.items():
            zf.writestr(name, data)
    return buf.getvalue()


def _import(ctx, headers, blob, **form):
    return ctx.client.post('/api/backup/import', headers=headers, data=form,
                           files={'file': ('b.zip', blob, 'application/zip')})


def test_backup_roundtrip(ctx):
    a = ctx.admin()
    url = ctx.client.post('/api/logo/upload', headers=a,
                          files={'file': ('a.png', png_bytes(), 'image/png')}).json()['url']
    ctx.client.post('/api/clients', headers=a, json={'company': 'Keep', 'logo_url': url})
    blob = ctx.client.get('/api/backup/export', headers=a).content
    names = zipfile.ZipFile(io.BytesIO(blob)).namelist()
    assert {'manifest.json', 'data.json', 'users.json', url} <= set(names)
    assert '.jwt_secret' not in ' '.join(names)

    ctx.client.post('/api/clients', headers=a, json={'company': 'Later'})
    (ctx.data_dir / url).unlink()
    r = _import(ctx, a, blob)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body['records'] == 1 and body['logos'] == 1 and body['users'] is False
    assert body['snapshot'] and (ctx.data_dir / 'backups' / body['snapshot']).exists()
    assert (ctx.data_dir / url).exists()
    assert [x['company'] for x in ctx.client.get('/api/clients', headers=a).json()['records']] == ['Keep']


def test_backup_does_not_restore_users_unless_asked(ctx):
    a = ctx.admin()
    blob = ctx.client.get('/api/backup/export', headers=a).content
    ctx.client.post('/api/users', headers=a, json={'username': 'newbie', 'password': 'newbie-password'})
    _import(ctx, a, blob)
    assert 'newbie' in [u['username'] for u in ctx.client.get('/api/users', headers=a).json()['users']]
    r = _import(ctx, a, blob, restore_users='true')
    assert r.json()['users'] is True
    ctx.client.cookies.clear()
    assert 'newbie' not in [u['username'] for u in ctx.client.get('/api/users', headers=a).json()['users']]


def test_backup_users_without_admin_is_refused_and_nothing_changes(ctx):
    a = ctx.admin()
    ctx.client.post('/api/clients', headers=a, json={'company': 'Original'})
    users = {'users': [{'username': 'v', 'password_hash': '$2b$12$abc', 'role': 'viewer'}]}
    blob = _zip({'data.json': json.dumps({'records': []}), 'users.json': json.dumps(users),
                 'logos/new.png': png_bytes()})
    r = _import(ctx, a, blob, restore_users='true')
    assert r.status_code == 400 and 'admin' in r.json()['detail']
    assert ctx.client.get('/api/clients', headers=a).json()['total'] == 1
    assert not (ctx.data_dir / 'logos' / 'new.png').exists()


@pytest.mark.parametrize('bad_name', ['../evil.png', '/abs.png', 'logos/../../x.png', 'C:/x.png'])
def test_backup_rejects_path_traversal(ctx, bad_name):
    blob = _zip({'data.json': json.dumps({'records': []}), bad_name: b'x'})
    assert _import(ctx, ctx.admin(), blob).status_code == 400


def test_backup_skips_unsafe_logo_entries(ctx):
    blob = _zip({'data.json': json.dumps({'records': []}),
                 'logos/ok.png': png_bytes(),
                 'logos/script.js': b'alert(1)',
                 'logos/fake.png': b'<html>not an image</html>',
                 'logos/.hidden.png': png_bytes()})
    body = _import(ctx, ctx.admin(), blob).json()
    assert body['logos'] == 1
    assert sorted(body['skipped']) == ['logos/.hidden.png', 'logos/fake.png', 'logos/script.js']
    assert not (ctx.data_dir / 'logos' / 'script.js').exists()


def test_backup_zip_bomb_rejected(make_ctx):
    c = make_ctx(seed_records=[], BACKUP_MAX_UNCOMPRESSED_MB='1')
    blob = _zip({'data.json': json.dumps({'records': []}), 'logos/a.png': b'\x00' * (3 * 1024 * 1024)})
    assert len(blob) < 100 * 1024  # tiny on the wire
    r = _import(c, c.admin(), blob)
    assert r.status_code == 400 and 'uncompressed' in r.json()['detail']


def test_backup_invalid_data_json(ctx):
    a = ctx.admin()
    assert _import(ctx, a, _zip({'data.json': '{"records": 5}'})).status_code == 400
    assert _import(ctx, a, _zip({'x.txt': 'hi'})).status_code == 400
    assert _import(ctx, a, b'not a zip').status_code == 400


def test_auto_snapshots_are_pruned(make_ctx):
    c = make_ctx(seed_records=[], AUTO_SNAPSHOTS_KEEP='2')
    a = c.admin()
    blob = _zip({'data.json': json.dumps({'records': []})})
    for _ in range(4):
        assert _import(c, a, blob).status_code == 200
    assert len(list((c.data_dir / 'backups').glob('auto-*.zip'))) == 2


# ---- Excel ----------------------------------------------------------------------
def _xlsx(rows):
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_excel_import_by_header_and_types(ctx):
    a = ctx.admin()
    blob = _xlsx([
        ['公司', '办公室', '部门', '负责人', '合作时间'],
        ['腾讯科技', 'SZH', '审计;税务', '王五', datetime(2023, 6, 1)],
        ['Foo Ltd', 'bji', None, '赵六；钱七', 2021],
        [None, None, None, None, None],
        ['Bar', None, None, None, '2022.07'],
    ])
    r = ctx.client.post('/api/import-excel', headers=a, files={'file': ('c.xlsx', blob, 'application/octet-stream')})
    assert r.status_code == 200, r.text
    assert r.json()['added'] == 3 and r.json()['logo_matched'] == 1
    recs = {x['company']: x for x in ctx.client.get('/api/clients', headers=a).json()['records']}
    assert recs['腾讯科技']['cooperation_date'] == '2023-06-01'
    assert recs['腾讯科技']['departments'] == ['审计', '税务']
    assert recs['腾讯科技']['logo_url'].startswith('https://')
    assert recs['Foo Ltd']['office_city'] == '北京' and recs['Foo Ltd']['cooperation_date'] == '2021'
    assert recs['Foo Ltd']['owners'] == ['赵六', '钱七']
    assert recs['Bar']['cooperation_date'] == '2022-07'


def test_excel_import_positional_columns(ctx):
    a = ctx.admin()
    blob = _xlsx([['A', 'B', 'C', 'D', 'E'], ['Pos Co', 'SHA', '咨询', '甲', '2020-01-02']])
    ctx.client.post('/api/import-excel', headers=a, files={'file': ('c.xlsx', blob)})
    rec = ctx.client.get('/api/clients', headers=a).json()['records'][0]
    assert (rec['company'], rec['office_city'], rec['departments'], rec['owners'], rec['cooperation_date']) == \
        ('Pos Co', '上海', ['咨询'], ['甲'], '2020-01-02')


def test_excel_import_rejects_garbage(ctx):
    r = ctx.client.post('/api/import-excel', headers=ctx.admin(), files={'file': ('c.xlsx', b'garbage')})
    assert r.status_code == 400


def test_excel_export_roundtrip(ctx):
    a = ctx.admin()
    ctx.client.post('/api/clients', headers=a, json={
        'company': 'Exp', 'office_code': 'GZH', 'departments': '审计;税务', 'owners': 'A',
        'cooperation_date': '2024-01-01'})
    blob = ctx.client.get('/api/export-excel', headers=a).content
    ws = load_workbook(io.BytesIO(blob)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == '租客/买方'
    assert rows[1] == ('Exp', 'GZH', '华南', '审计；税务', 'A', '2024-01-01')
    # and it imports back cleanly
    ctx.client.post('/api/import-excel', headers=a, files={'file': ('e.xlsx', blob)})
    recs = ctx.client.get('/api/clients', headers=a).json()['records']
    assert recs[0]['departments'] == recs[1]['departments'] == ['审计', '税务']


def test_brand_keyword_override(ctx):
    (ctx.data_dir / 'brand_keywords.json').write_text(json.dumps({'Acme': 'https://cdn.example/acme.png', '腾讯': ''}),
                                                     encoding='utf-8')
    from logowall.constants import match_brand_logo
    assert match_brand_logo('Acme Corp') == ('Acme', 'https://cdn.example/acme.png')
    assert match_brand_logo('腾讯科技') == (None, None)
